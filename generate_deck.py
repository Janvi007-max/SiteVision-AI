"""
SiteVision AI - Automated PowerPoint Presentation Generator
Creates a sleek, 16:9 widescreen, dark-themed presentation deck.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="SiteVision_AI_Presentation.pptx"):
    prs = Presentation()
    # Set 16:9 Widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette (Dark Theme matching 21st.dev)
    BG_DARK = RGBColor(14, 14, 18)       # #0e0e12
    CARD_BG = RGBColor(24, 24, 30)       # #18181e
    TEXT_WHITE = RGBColor(244, 244, 245) # #f4f4f5
    TEXT_MUTED = RGBColor(161, 161, 170) # #a1a1aa
    ACCENT_CYAN = RGBColor(6, 182, 212)  # #06b6d4
    ACCENT_GREEN = RGBColor(16, 185, 129)# #10b981
    ACCENT_RED = RGBColor(239, 68, 68)   # #ef4444
    ACCENT_AMBER = RGBColor(245, 158, 11)# #f59e0b

    def add_blank_slide():
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return slide

    def add_header(slide, title_text, category_text="SITEVISION AI // ARCHITECTURE"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.color.rgb = RGBColor(45, 45, 55)
            card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = add_blank_slide()
    # Title badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(16, 185, 129)
    badge.line.fill.background()
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "● AUTONOMOUS SAFETY SYSTEM"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(10, 30, 20)
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.8))
    tf1 = t_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "SiteVision AI"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Autonomous Construction Safety & DLP Access Control Web App"
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT_CYAN

    # Subtitle / Details Card
    add_card(s1, 0.8, 4.6, 11.7, 1.8)
    d_box = s1.shapes.add_textbox(Inches(1.1), Inches(4.8), Inches(11.1), Inches(1.4))
    tf_d = d_box.text_frame
    p_d1 = tf_d.paragraphs[0]
    p_d1.text = "Key Innovations:"
    p_d1.font.bold = True
    p_d1.font.size = Pt(13)
    p_d1.font.color.rgb = TEXT_WHITE

    items = [
        "Real-Time PPE Detection (Hardhats, High-Vis Safety Vests) via YOLOv8 + OpenCV",
        "Zero-Tolerance DLP Security Rule: Smoking / Open Flame -> Immediate ID Suspension",
        "AI Safety Officer Audit (Google Gemini LLM + Offline NLP) & Web Speech Voice Broadcast",
        "1-Click OSHA Incident PDF Report Generator with Full Inspection Telemetry"
    ]
    for item in items:
        p = tf_d.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: Problem Statement & Industry Need
    # =========================================================================
    s2 = add_blank_slide()
    add_header(s2, "The Problem: Critical Safety & Compliance Hazards on Site", "INDUSTRY CHALLENGE")

    add_card(s2, 0.8, 1.8, 3.6, 4.8)
    b1 = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.2), Inches(4.4))
    t1 = b1.text_frame
    t1.word_wrap = True
    t1.paragraphs[0].text = "1. Fatal Cranial & Visibility Risks"
    t1.paragraphs[0].font.bold = True
    t1.paragraphs[0].font.size = Pt(14)
    t1.paragraphs[0].font.color.rgb = ACCENT_AMBER
    p = t1.add_paragraph()
    p.text = "\n• Over 60% of construction head injuries occur due to workers failing to wear hardhats.\n\n• Low-visibility accidents in heavy machinery zones result in severe OSHA penalties and downtime."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    add_card(s2, 4.8, 1.8, 3.6, 4.8, border_color=ACCENT_RED)
    b2 = s2.shapes.add_textbox(Inches(5.0), Inches(2.0), Inches(3.2), Inches(4.4))
    t2 = b2.text_frame
    t2.word_wrap = True
    t2.paragraphs[0].text = "2. Fire & Contraband DLP Hazards"
    t2.paragraphs[0].font.bold = True
    t2.paragraphs[0].font.size = Pt(14)
    t2.paragraphs[0].font.color.rgb = ACCENT_RED
    p = t2.add_paragraph()
    p.text = "\n• Unauthorized smoking near combustible materials / scaffolds is a leading cause of job site fires.\n\n• Existing access gates lack real-time contraband detection and automated badge suspension."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    add_card(s2, 8.8, 1.8, 3.7, 4.8)
    b3 = s2.shapes.add_textbox(Inches(9.0), Inches(2.0), Inches(3.3), Inches(4.4))
    t3 = b3.text_frame
    t3.word_wrap = True
    t3.paragraphs[0].text = "3. Manual Inspection Bottlenecks"
    t3.paragraphs[0].font.bold = True
    t3.paragraphs[0].font.size = Pt(14)
    t3.paragraphs[0].font.color.rgb = ACCENT_CYAN
    p = t3.add_paragraph()
    p.text = "\n• Human safety officers cannot monitor hundreds of entry points simultaneously.\n\n• Manual paper logging is slow, error-prone, and provides no immediate physical gate enforcement."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Proposed Solution - SiteVision AI
    # =========================================================================
    s3 = add_blank_slide()
    add_header(s3, "The Solution: Autonomous Vision & DLP Access Governance", "SOLUTION OVERVIEW")

    add_card(s3, 0.8, 1.8, 5.6, 4.8)
    b_sol1 = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.4))
    t_sol1 = b_sol1.text_frame
    t_sol1.word_wrap = True
    t_sol1.paragraphs[0].text = "🛡️ Core Vision & Governance Capabilities"
    t_sol1.paragraphs[0].font.bold = True
    t_sol1.paragraphs[0].font.size = Pt(15)
    t_sol1.paragraphs[0].font.color.rgb = ACCENT_GREEN

    points_sol1 = [
        ("Multi-Layer PPE Verification", "Simultaneously scans headgear color profiles and retro-reflective torso bands in milliseconds."),
        ("Automated DLP Policy Enforcement", "Instantly flags smoking/open flame and revokes Worker ID credentials before site entry."),
        ("Multi-Modal Alerting", "Broadcasts high-clarity voice alerts via Web Speech API and triggers audible klaxon alarms."),
        ("OSHA Incident Reporting", "Generates verifiable PDF audit certificates with embedded visual proofs in 1-click.")
    ]
    for title, desc in points_sol1:
        p_t = t_sol1.add_paragraph()
        p_t.text = f"\n• {title}:"
        p_t.font.bold = True
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = TEXT_WHITE
        p_d = t_sol1.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_MUTED

    # Right side: Architecture flow cards
    flows = [
        ("1. Real-Time Vision Input", "Webcam Feed, On-Site Uploads, or 1-Click Demo Scenarios", ACCENT_CYAN),
        ("2. Dual-Engine Processing", "YOLOv8 Object Detection + OpenCV Color/Spatial Heuristics", ACCENT_AMBER),
        ("3. AI Safety Officer Audit", "Google Gemini LLM + Built-in Local NLP Safety Engine", ACCENT_GREEN),
        ("4. Automated Action & Export", "Gate Lock / Clear, Audio Broadcast, and OSHA PDF Log", ACCENT_RED)
    ]
    for i, (f_title, f_desc, col) in enumerate(flows):
        top_pos = 1.8 + i * 1.25
        add_card(s3, 6.8, top_pos, 5.7, 1.1, border_color=col)
        f_box = s3.shapes.add_textbox(Inches(7.0), Inches(top_pos + 0.1), Inches(5.3), Inches(0.9))
        tf_f = f_box.text_frame
        p_ft = tf_f.paragraphs[0]
        p_ft.text = f_title
        p_ft.font.bold = True
        p_ft.font.size = Pt(12)
        p_ft.font.color.rgb = col
        p_fd = tf_f.add_paragraph()
        p_fd.text = f_desc
        p_fd.font.size = Pt(10)
        p_fd.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: System Architecture & Workflow Pipeline
    # =========================================================================
    s4 = add_blank_slide()
    add_header(s4, "End-to-End System Pipeline & Data Flow", "SYSTEM ARCHITECTURE")

    pipeline_steps = [
        ("1. INGESTION", "• Camera / Upload\n• RGB normalization\n• Frame pre-processing", ACCENT_CYAN),
        ("2. DETECTION", "• YOLOv8 Person ROI\n• Head dome crop\n• Torso HSV analysis", ACCENT_AMBER),
        ("3. DLP AUDIT", "• Cigarette / Flame scan\n• ID suspension check\n• Status evaluation", ACCENT_RED),
        ("4. LLM OFFICER", "• Gemini 2.5 Flash\n• Local NLP Fallback\n• Natural language audit", ACCENT_GREEN),
        ("5. ACTION & UI", "• 21st.dev Dark UI\n• Web Speech TTS\n• OSHA PDF Export", ACCENT_CYAN)
    ]
    for i, (p_name, p_body, col) in enumerate(pipeline_steps):
        left_pos = 0.8 + i * 2.45
        add_card(s4, left_pos, 1.8, 2.3, 4.8, border_color=col)
        p_box = s4.shapes.add_textbox(Inches(left_pos + 0.15), Inches(2.0), Inches(2.0), Inches(4.4))
        tf_p = p_box.text_frame
        tf_p.word_wrap = True
        p_hd = tf_p.paragraphs[0]
        p_hd.text = p_name
        p_hd.font.bold = True
        p_hd.font.size = Pt(13)
        p_hd.font.color.rgb = col
        p_bd = tf_p.add_paragraph()
        p_bd.text = f"\n{p_body}"
        p_bd.font.size = Pt(10.5)
        p_bd.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: DLP Policy Matrix & Access Clearance
    # =========================================================================
    s5 = add_blank_slide()
    add_header(s5, "DLP Rule Enforcement Matrix: Zero-Tolerance Policy", "ACCESS GOVERNANCE")

    matrix_cards = [
        ("🟢 SAFE STATUS", "100% PPE Clearance", "• Hardhat Detected (HSV Verified)\n• Safety Vest Detected (Class-2)\n• Zero Contraband / Smoking", "Turnstile 01 UNLOCKED\nAccess Token Issued\nVoice: 'Access Granted'", ACCENT_GREEN),
        ("🟡 WARNING STATUS", "PPE Non-Compliance", "• Missing Hardhat OR Missing Vest\n• High head/torso injury risk\n• Non-compliant with OSHA 1926", "Turnstile HELD\nRedirect to Staging Area\nVoice: 'Equip Required Gear'", ACCENT_AMBER),
        ("🔴 ID BLOCKED (DLP)", "Critical Security Violation", "• Cigarette / Flame Detected\n• Active fire & explosion hazard\n• Severe safety infraction", "Turnstile LOCKED\nWorker ID SUSPENDED\nAudio: Klaxon Alarm + TTS", ACCENT_RED)
    ]
    for i, (m_title, m_sub, m_rules, m_action, col) in enumerate(matrix_cards):
        left_pos = 0.8 + i * 4.0
        add_card(s5, left_pos, 1.8, 3.7, 4.8, border_color=col)
        m_box = s5.shapes.add_textbox(Inches(left_pos + 0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf_m = m_box.text_frame
        tf_m.word_wrap = True

        p_m1 = tf_m.paragraphs[0]
        p_m1.text = m_title
        p_m1.font.bold = True
        p_m1.font.size = Pt(14)
        p_m1.font.color.rgb = col

        p_m2 = tf_m.add_paragraph()
        p_m2.text = m_sub
        p_m2.font.size = Pt(11)
        p_m2.font.color.rgb = TEXT_WHITE

        p_m3 = tf_m.add_paragraph()
        p_m3.text = f"\nConditions:\n{m_rules}"
        p_m3.font.size = Pt(10)
        p_m3.font.color.rgb = TEXT_MUTED

        p_m4 = tf_m.add_paragraph()
        p_m4.text = f"\nEnforcement Action:\n{m_action}"
        p_m4.font.bold = True
        p_m4.font.size = Pt(10)
        p_m4.font.color.rgb = col

    # =========================================================================
    # SLIDE 6: AI Safety Officer & LLM Auditing
    # =========================================================================
    s6 = add_blank_slide()
    add_header(s6, "Autonomous AI Safety Officer (LLM + NLP Engine)", "NLP & LLM INTEGRATION")

    add_card(s6, 0.8, 1.8, 5.6, 4.8)
    b_llm1 = s6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_l1 = b_llm1.text_frame
    tf_l1.word_wrap = True
    tf_l1.paragraphs[0].text = "🧠 Dual-Mode Intelligence Architecture"
    tf_l1.paragraphs[0].font.bold = True
    tf_l1.paragraphs[0].font.size = Pt(14)
    tf_l1.paragraphs[0].font.color.rgb = ACCENT_CYAN

    l_items = [
        ("Online Mode (Google Gemini LLM)", "Uses Gemini 2.5 Flash API to generate authoritative 2-3 sentence construction safety debriefs contextualized to exact site hazards."),
        ("Offline Mode (Local Zero-Latency NLP)", "Built-in heuristic NLP engine ensures 100% uptime and instant evaluations even with zero internet connection or API keys."),
        ("OSHA Regulatory Alignment", "Automatically links specific detections to OSHA 1926.151 (Fire Prevention) & 1926.100 (Head Protection) standards.")
    ]
    for title, desc in l_items:
        p_t = tf_l1.add_paragraph()
        p_t.text = f"\n• {title}:"
        p_t.font.bold = True
        p_t.font.size = Pt(11)
        p_t.font.color.rgb = TEXT_WHITE
        p_d = tf_l1.add_paragraph()
        p_d.text = f"  {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    # Right side: Sample LLM Outputs
    add_card(s6, 6.8, 1.8, 5.7, 4.8)
    b_llm2 = s6.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tf_l2 = b_llm2.text_frame
    tf_l2.word_wrap = True
    tf_l2.paragraphs[0].text = "📝 Real-Time AI Audit Debrief Samples"
    tf_l2.paragraphs[0].font.bold = True
    tf_l2.paragraphs[0].font.size = Pt(14)
    tf_l2.paragraphs[0].font.color.rgb = ACCENT_GREEN

    samples = [
        ("DLP Block Output", "🚨 CRITICAL DLP BREACH: Worker observed smoking on active job site. Per OSHA 1926.151 fire prevention protocol, Worker ID badge clearance is immediately BLOCKED and site turnstiles locked.", ACCENT_RED),
        ("Warning Output", "⚠️ PPE WARNING: High-visibility vest verified, but mandatory hardhat is missing. Personnel must equip approved head protection before proceeding past staging.", ACCENT_AMBER),
        ("Safe Clearance Output", "✅ SAFETY VERIFIED: Personnel successfully cleared with approved hardhat and Class-2 high-visibility vest. Site access is granted under standard protocol.", ACCENT_GREEN)
    ]
    for s_title, s_text, col in samples:
        p_st = tf_l2.add_paragraph()
        p_st.text = f"\n{s_title}:"
        p_st.font.bold = True
        p_st.font.size = Pt(10.5)
        p_st.font.color.rgb = col
        p_sb = tf_l2.add_paragraph()
        p_sb.text = f'"{s_text}"'
        p_sb.font.size = Pt(9.5)
        p_sb.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 7: High-Impact Features: Voice Alerts & OSHA PDF
    # =========================================================================
    s7 = add_blank_slide()
    add_header(s7, "Feature Highlights: Voice Broadcast & 1-Click OSHA PDF", "INNOVATIVE EXTENSIONS")

    add_card(s7, 0.8, 1.8, 5.6, 4.8, border_color=ACCENT_CYAN)
    b_v = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_v = b_v.text_frame
    tf_v.word_wrap = True
    tf_v.paragraphs[0].text = "🔊 Live Web Speech Audio & Klaxon Alert"
    tf_v.paragraphs[0].font.bold = True
    tf_v.paragraphs[0].font.size = Pt(14)
    tf_v.paragraphs[0].font.color.rgb = ACCENT_CYAN
    p_v = tf_v.add_paragraph()
    p_v.text = "\n• Zero External Dependencies: Utilizes native browser Web Speech API (`speechSynthesis`) and Web Audio oscillators.\n\n• Spoken Voice Announcements: Hands-free vocal broadcasting for security personnel and workers.\n\n• Audio Klaxon: Emits an immediate audible siren pulse upon smoking / open flame detection.\n\n• Fully Toggleable: Can be enabled/disabled instantly from the sidebar."
    p_v.font.size = Pt(11)
    p_v.font.color.rgb = TEXT_MUTED

    add_card(s7, 6.8, 1.8, 5.7, 4.8, border_color=ACCENT_GREEN)
    b_pdf = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tf_pdf = b_pdf.text_frame
    tf_pdf.word_wrap = True
    tf_pdf.paragraphs[0].text = "📄 1-Click OSHA Incident PDF Generator"
    tf_pdf.paragraphs[0].font.bold = True
    tf_pdf.paragraphs[0].font.size = Pt(14)
    tf_pdf.paragraphs[0].font.color.rgb = ACCENT_GREEN
    p_p = tf_pdf.add_paragraph()
    p_p.text = "\n• Instant Audit Export: Generates professional single-page A4 audit certificates using lightweight `fpdf2`.\n\n• Structured Metadata: Includes Worker Token ID, Zone, Clearance Status, and Timestamp.\n\n• Visual Evidence Embedding: Renders high-resolution camera snapshots with glowing HUD bounding boxes.\n\n• Official Sign-Off: Digital verification line compliant with OSHA Standard 1926."
    p_p.font.size = Pt(11)
    p_p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 8: Tech Stack & Architecture Highlights
    # =========================================================================
    s8 = add_blank_slide()
    add_header(s8, "Tech Stack & Engineering Design Principles", "TECHNICAL FOUNDATION")

    techs = [
        ("Streamlit + 21st.dev CSS", "Frontend & UI", "Dark zinc `#09090b` palette, glassmorphism, responsive 2-column inspection layout, glowing status cards.", ACCENT_CYAN),
        ("Ultralytics YOLOv8", "Deep Learning Object Detection", "Lightweight `yolov8n` nano architecture for high-speed person and contraband localization.", ACCENT_AMBER),
        ("OpenCV Computer Vision", "Heuristic Spatial Engine", "HSV color segmentation for hardhats and reflective vest bands, plus HUD bounding box rendering.", ACCENT_GREEN),
        ("Google Gemini & Local NLP", "Language & Reasoning Engine", "Multi-modal contextual audit reports with instantaneous offline template fallback.", ACCENT_CYAN),
        ("fpdf2 PDF Exporter", "Reporting & Compliance", "Pure Python lightweight A4 document generator with zero heavy C-dependencies.", ACCENT_RED)
    ]
    for i, (t_name, t_cat, t_desc, col) in enumerate(techs):
        top_pos = 1.8 + i * 1.0
        add_card(s8, 0.8, top_pos, 11.7, 0.88, border_color=col)
        t_box = s8.shapes.add_textbox(Inches(1.0), Inches(top_pos + 0.05), Inches(11.3), Inches(0.8))
        tf_t = t_box.text_frame
        p_t1 = tf_t.paragraphs[0]
        p_t1.text = f"{t_name}  —  [{t_cat}]"
        p_t1.font.bold = True
        p_t1.font.size = Pt(11)
        p_t1.font.color.rgb = col
        p_t2 = tf_t.add_paragraph()
        p_t2.text = t_desc
        p_t2.font.size = Pt(10)
        p_t2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Project Scalability & Future Roadmap
    # =========================================================================
    s9 = add_blank_slide()
    add_header(s9, "Future Roadmap: Enterprise Scaling & Next Steps", "FUTURE VISION")

    roadmap = [
        ("Phase 1: Multi-Camera RTSP Streams", "Integrate continuous 30 FPS CCTV stream analysis with multi-worker ByteTrack / DeepSORT tracking across active site sectors.", ACCENT_CYAN),
        ("Phase 2: Virtual Danger Zone Geofencing", "Allow safety supervisors to draw virtual red-lines on camera feeds (e.g. under crane booms) triggering instant perimeter breach alarms.", ACCENT_AMBER),
        ("Phase 3: IoT RFID Turnstile Hardware API", "Connect direct Webhook / MQTT triggers to physical site turnstiles and strobe lights to physically lock out violators in real-time.", ACCENT_RED),
        ("Phase 4: Site Safety Analytics Dashboard", "Aggregate multi-day compliance trends, site safety index scores (0-100%), and automatic subcontractor audit reports.", ACCENT_GREEN)
    ]
    for i, (r_title, r_desc, col) in enumerate(roadmap):
        top_pos = 1.8 + i * 1.25
        add_card(s9, 0.8, top_pos, 11.7, 1.1, border_color=col)
        r_box = s9.shapes.add_textbox(Inches(1.1), Inches(top_pos + 0.1), Inches(11.1), Inches(0.9))
        tf_r = r_box.text_frame
        p_rt = tf_r.paragraphs[0]
        p_rt.text = r_title
        p_rt.font.bold = True
        p_rt.font.size = Pt(13)
        p_rt.font.color.rgb = col
        p_rd = tf_r.add_paragraph()
        p_rd.text = r_desc
        p_rd.font.size = Pt(10.5)
        p_rd.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 10: Conclusion & Summary
    # =========================================================================
    s10 = add_blank_slide()
    # Title badge
    badge10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(2.8), Inches(0.45))
    badge10.fill.solid()
    badge10.fill.fore_color.rgb = ACCENT_GREEN
    badge10.line.fill.background()
    p_b10 = badge10.text_frame.paragraphs[0]
    p_b10.text = "● PROJECT SUMMARY"
    p_b10.font.size = Pt(10)
    p_b10.font.bold = True
    p_b10.font.color.rgb = RGBColor(10, 30, 20)
    p_b10.alignment = PP_ALIGN.CENTER

    t_box10 = s10.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.5), Inches(1.2))
    tf10 = t_box10.text_frame
    p10 = tf10.paragraphs[0]
    p10.text = "SiteVision AI: Smarter, Safer Construction Sites"
    p10.font.size = Pt(32)
    p10.font.bold = True
    p10.font.color.rgb = TEXT_WHITE

    add_card(s10, 0.8, 3.5, 11.7, 3.2)
    c_box = s10.shapes.add_textbox(Inches(1.1), Inches(3.7), Inches(11.1), Inches(2.8))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True

    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "Key Takeaways:"
    p_c1.font.bold = True
    p_c1.font.size = Pt(14)
    p_c1.font.color.rgb = ACCENT_CYAN

    concl_points = [
        "100% Autonomous PPE & Contraband Detection: Eliminates human error and automates turnstile access control.",
        "Zero-Tolerance DLP Smoking Rule: Prevents catastrophic job site fire hazards through instant badge suspension.",
        "Production-Grade Dark UI & Voice AI: Elevates standard computer vision into an immersive, industry-level safety system.",
        "OSHA-Compliant Audit Logging: 1-click PDF certificates streamline site compliance and regulatory reviews."
    ]
    for pt in concl_points:
        p = tf_c.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE

    prs.save(output_path)
    print(f"[SiteVision AI] Presentation successfully saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
